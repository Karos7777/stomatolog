"""Скачать официальные реестры Минздрава КР и сохранить в data/registry/.

    python -m tools.license_registry

Берёт со страницы https://med.kg/lisenzirovanie:
  • «Реестр выданных лицензий на медицинскую деятельность» (xlsx)
  • «Список субъектов предпринимательства, осуществляющих безлицензионную деятельность» (docx)
Нужны пакеты openpyxl и python-docx (только для этого скрипта).
"""
import html
import io
import json
import re
import sys
import urllib.parse
import urllib.request
from datetime import date
from pathlib import Path
from typing import Dict, List, Tuple

PAGE = "https://med.kg/lisenzirovanie?locale=ru"
OUT_DIR = Path(__file__).resolve().parent.parent / "data" / "registry"


def _get(url: str, timeout: int = 90) -> bytes:
    with urllib.request.urlopen(url, timeout=timeout) as resp:
        return resp.read()


def _title(href: str) -> str:
    return re.sub(r"^[0-9a-f\-]{36}-", "", Path(href).stem).strip()


def find_links(page_html: str) -> Dict[str, Tuple[str, str]]:
    """{'licenses': (url, title), 'unlicensed': (url, title)} по тексту и расширению ссылок."""
    found: Dict[str, Tuple[str, str]] = {}
    for m in re.finditer(r'href="([^"]+)"', page_html):
        href = html.unescape(m.group(1))
        low = href.lower()
        absolute = urllib.parse.urljoin("https://med.kg/", urllib.parse.quote(href, safe="/:?=&%"))
        if low.endswith(".xlsx") and "медицинскую деятельность" in low and "licenses" not in found:
            found["licenses"] = (absolute, _title(href))
        elif low.endswith(".docx") and "без лицензион" in low and "unlicensed" not in found:
            found["unlicensed"] = (absolute, _title(href))
    return found


def _clean(v) -> str:
    return re.sub(r"\s+", " ", str(v).replace("_x000D_", " ")).strip() if v is not None else ""


def parse_licenses(xlsx: bytes) -> List[Dict]:
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(xlsx), read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    out = []
    for r in rows[1:]:
        if not any(r):
            continue
        no, number, name, activity, region, district, locality, address, issued = (list(r) + [None] * 9)[:9]
        out.append({"number": _clean(number), "name": _clean(name), "activity": _clean(activity),
                    "region": _clean(region), "address": _clean(address),
                    "issued": _clean(issued)[:10], "dental": "стомат" in _clean(activity).lower()})
    return out


def parse_unlicensed(docx_bytes: bytes) -> List[Dict]:
    import docx
    d = docx.Document(io.BytesIO(docx_bytes))
    out = []
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) < 7 or not cells[0].isdigit():
                continue
            out.append({"name": cells[1], "activity": cells[2], "region": cells[3],
                        "address": cells[6], "dental": "стомат" in cells[2].lower()})
    return out


def main() -> int:
    page = _get(PAGE).decode("utf-8", errors="replace")
    links = find_links(page)
    if "licenses" not in links:
        print("Не нашёл ссылку на реестр лицензий на странице", PAGE)
        return 1
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    today = date.today().isoformat()

    lic_url, lic_title = links["licenses"]
    licenses = parse_licenses(_get(lic_url))
    bishkek = [r for r in licenses if "бишкек" in r["region"].lower()]
    as_of = re.search(r"на (\d{1,2}\.\d{1,2}\.\s?\d{4})", lic_title)
    (OUT_DIR / "medical_licenses_bishkek.json").write_text(json.dumps({
        "meta": {"source_page": PAGE, "source_file": lic_url, "title": lic_title,
                 "as_of": as_of.group(1).replace(" ", "") if as_of else None, "downloaded": today,
                 "records_total": len(licenses), "records_bishkek": len(bishkek)},
        "records": bishkek}, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"Лицензии: {len(licenses)} всего, {len(bishkek)} в Бишкеке")

    if "unlicensed" in links:
        un_url, un_title = links["unlicensed"]
        unlicensed = parse_unlicensed(_get(un_url))
        (OUT_DIR / "unlicensed.json").write_text(json.dumps({
            "meta": {"source_page": PAGE, "source_file": un_url, "title": un_title, "downloaded": today,
                     "records": len(unlicensed)},
            "records": unlicensed}, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
        print(f"Без лицензии: {len(unlicensed)} записей")
    return 0


if __name__ == "__main__":
    sys.exit(main())
